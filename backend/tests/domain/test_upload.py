"""Reading uploaded files: Excel workbooks, text encodings, and the limits that keep a hostile
file from costing the server anything. Every row in here is invented."""

import io
import time
import zipfile
from datetime import date, datetime

import openpyxl
import pytest
from openpyxl import Workbook

from app.domain.csv_import import HEADER_SCAN_ROWS, MAX_COLUMNS, MAX_ROWS, CsvImportError, parse_csv
from app.domain.upload import (
    MAX_SHEETS,
    MAX_WORKBOOK_UNCOMPRESSED,
    decode_text,
    file_kind,
    parse_upload,
    read_workbook,
)

HEADER = ["Type", "Date", "Item", "Amount (RM)", "Category"]
TABLE = [
    HEADER,
    ["expense", datetime(2025, 8, 1), " Kopi Corner", 6.8, "Drink / Dessert"],
    ["expense", datetime(2025, 8, 2), "Rice Stall", 15.9, " Food"],
    ["income", datetime(2025, 8, 3), "Monthly Allowance", 1000, "Allowance"],
]


def workbook_bytes(sheets: dict[str, list[list]], hidden: tuple[str, ...] = ()) -> bytes:
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
        if name in hidden:
            sheet.sheet_state = "hidden"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def excel(rows: list[list]) -> bytes:
    return workbook_bytes({"Sheet1": rows})


def summary(preview):
    return [(r.date, r.amount_sen, r.type, r.category_guess, r.description) for r in preview.rows]


# ---------------------------------------------------------------- which files are accepted
@pytest.mark.parametrize(
    ("name", "kind"),
    [
        ("bank.csv", "text"),
        ("BANK.CSV", "text"),
        ("bank.tsv", "text"),
        ("bank.txt", "text"),
        ("budget.xlsx", "excel"),
        ("Budget.XLSM", "excel"),
    ],
)
def test_supported_file_names(name, kind):
    assert file_kind(name) == kind


@pytest.mark.parametrize("name", ["bank.pdf", "bank", "", None, "notes.docx", "data.xlsb"])
def test_other_file_names_are_refused(name):
    with pytest.raises(CsvImportError) as err:
        file_kind(name)

    assert err.value.code == "WRONG_FILE_TYPE"
    assert ".xlsx" in err.value.message


def test_an_old_xls_file_is_refused_with_the_way_to_convert_it():
    with pytest.raises(CsvImportError) as err:
        file_kind("budget.xls")

    assert err.value.code == "WRONG_FILE_TYPE"
    assert "Save As" in err.value.message and ".xlsx" in err.value.message


# ---------------------------------------------------------------- workbooks read like the CSV
def test_an_xlsx_table_gives_the_same_rows_as_the_same_table_as_csv():
    csv_text = (
        "Type,Date,Item,Amount (RM),Category\n"
        "expense,1/8/2025, Kopi Corner,6.8,Drink / Dessert\n"
        "expense,2/8/2025,Rice Stall,15.9, Food\n"
        "income,3/8/2025,Monthly Allowance,1000,Allowance\n"
    )

    from_excel = parse_upload("excel", excel(TABLE))

    assert summary(from_excel) == summary(parse_csv(csv_text))
    assert [r.type for r in from_excel.rows] == ["expense", "expense", "income"]
    assert all(r.issues == () for r in from_excel.rows)
    assert from_excel.notes == ()


def test_amounts_that_excel_stored_with_binary_noise_are_read_as_what_it_showed():
    rows = [
        HEADER,
        ["expense", datetime(2025, 8, 1), "a", 12 * 1.1, "Food"],  # 13.200000000000001
        ["expense", datetime(2025, 8, 1), "b", 0.1 + 0.2, "Food"],  # 0.30000000000000004
        ["expense", datetime(2025, 8, 1), "c", 393.3, "Food"],
        ["expense", datetime(2025, 8, 1), "d", 1000, "Food"],
    ]

    preview = parse_upload("excel", excel(rows))

    assert [r.amount_sen for r in preview.rows] == [1320, 30, 39330, 100000]


def test_a_real_third_decimal_is_still_refused_not_rounded():
    rows = [HEADER, ["expense", datetime(2025, 8, 1), "a", 1.234, "Food"]]
    rows += [["expense", datetime(2025, 8, 1), "b", 2, "Food"]]

    preview = parse_upload("excel", excel(rows))

    assert [r.description for r in preview.rows] == ["b"]
    assert "amount" in preview.skipped[0].reason


def test_amounts_and_dates_stored_as_text_are_read_too():
    rows = [HEADER, ["expense", "1/8/2025", "a", "RM 6.80", "Food"]]
    rows += [["expense", datetime(2025, 8, 2, 14, 30), "b", "1,200.50", "Food"]]

    preview = parse_upload("excel", excel(rows))

    assert [(r.date, r.amount_sen) for r in preview.rows] == [
        (date(2025, 8, 1), 680),
        (date(2025, 8, 2), 120050),  # a date with a time of day is that day
    ]


def test_a_formula_with_no_stored_result_is_reported_not_guessed():
    rows = [HEADER, ["expense", datetime(2025, 8, 1), "a", "=5+1", "Food"]]
    rows += [["expense", datetime(2025, 8, 2), "b", 3, "Food"]]

    preview = parse_upload("excel", excel(rows))

    assert [r.description for r in preview.rows] == ["b"]
    assert "amount" in preview.skipped[0].reason


def test_title_lines_above_the_header_are_skipped_in_a_workbook():
    rows = [["My August budget"], [None, None], *TABLE]

    preview = parse_upload("excel", excel(rows))

    assert [r.row for r in preview.rows] == [1, 2, 3]


def test_rows_that_stop_before_the_last_column_are_still_read():
    rows = [HEADER, ["expense", datetime(2025, 8, 1), "Mamak dinner", 5]]  # no category cell

    (row,) = parse_upload("excel", excel(rows)).rows

    assert (row.amount_sen, row.category_guess) == (500, "Food")  # guessed: the file said nothing


def test_a_workbook_with_nothing_under_the_header_is_an_error():
    with pytest.raises(CsvImportError) as err:
        parse_upload("excel", excel([HEADER]))

    assert err.value.code == "CSV_NO_ROWS"


# ---------------------------------------------------------------- several sheets
def test_every_sheet_with_a_table_is_read_and_rows_stay_numbered_uniquely():
    july = [HEADER, ["expense", datetime(2025, 7, 1), "July lunch", 5, "Food"]]
    august = [HEADER, ["expense", datetime(2025, 8, 1), "August lunch", 6, "Food"]]
    august += [["expense", "not a date", "bad", 1, "Food"]]

    preview = parse_upload("excel", workbook_bytes({"Jul": july, "Aug": august}))

    assert [(r.row, r.description) for r in preview.rows] == [
        (1, "July lunch"),
        (2, "August lunch"),
    ]
    assert [(s.row, s.reason) for s in preview.skipped] == [(3, "sheet 'Aug': unparseable date")]


def test_a_sheet_that_is_not_a_table_is_named_in_a_note_and_the_rest_is_read():
    book = workbook_bytes({"Summary": [["Total spent"], [123]], "Aug": TABLE})

    preview = parse_upload("excel", book)

    assert len(preview.rows) == 3
    assert len(preview.notes) == 1 and "Summary" in preview.notes[0]


def test_a_hidden_sheet_is_not_read_and_says_so():
    book = workbook_bytes({"Aug": TABLE, "Old": TABLE}, hidden=("Old",))

    preview = parse_upload("excel", book)

    assert len(preview.rows) == 3
    assert any("Old" in note and "hidden" in note for note in preview.notes)


def test_a_workbook_with_no_table_anywhere_is_an_error():
    with pytest.raises(CsvImportError) as err:
        parse_upload("excel", workbook_bytes({"A": [["x", "y"], [1, 2]], "B": [["z"], [3]]}))

    assert err.value.code == "CSV_MISSING_COLUMNS"


def test_only_the_first_sheets_are_read_and_the_rest_are_mentioned():
    sheets = {f"S{i}": TABLE for i in range(MAX_SHEETS + 2)}

    preview = parse_upload("excel", workbook_bytes(sheets))

    assert len(preview.rows) == 3 * MAX_SHEETS
    assert any(str(MAX_SHEETS) in note for note in preview.notes)


# ---------------------------------------------------------------- limits
def test_a_stray_cell_far_below_the_data_does_not_make_the_server_scan_a_million_rows():
    book = Workbook()
    sheet = book.active
    for row in TABLE:
        sheet.append(row)
    sheet["A1048576"] = "stray"
    out = io.BytesIO()
    book.save(out)

    started = time.monotonic()
    preview = parse_upload("excel", out.getvalue())

    assert len(preview.rows) == 3
    assert time.monotonic() - started < 10


def test_too_many_rows_in_a_workbook_are_refused():
    rows = [HEADER] + [["expense", datetime(2025, 8, 1), "x", 1, "Food"]] * (MAX_ROWS + 1)

    with pytest.raises(CsvImportError) as err:
        parse_upload("excel", excel(rows))

    assert err.value.code == "CSV_TOO_MANY_ROWS"


def with_calc_chain(book: bytes, entries: int) -> bytes:
    """The workbook plus an `xl/calcChain.xml` like Excel writes for a sheet full of formulas: one
    `<c>` per formula cell, none of them inside a `<row>`."""
    chain = "".join(f'<c r="A{i}" i="1"/>' for i in range(1, entries + 1))
    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">{chain}'
        "</calcChain>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(book)) as source, zipfile.ZipFile(out, "w") as target:
        for part in source.infolist():
            target.writestr(part, source.read(part))
        target.writestr("xl/calcChain.xml", xml)
    return out.getvalue()


def test_a_workbook_with_many_formulas_is_not_mistaken_for_a_wide_sheet():
    # Found on the author's own workbook: Excel's formula chain lists every formula cell outside
    # any row, and counting those as one row refused the file as "more than 64 columns".
    book = with_calc_chain(excel(TABLE), entries=MAX_COLUMNS + 20)

    preview = parse_upload("excel", book)

    assert len(preview.rows) == len(TABLE) - 1


def test_too_many_columns_in_a_workbook_are_refused():
    wide = [[f"c{i}" for i in range(MAX_COLUMNS + 1)], [1] * (MAX_COLUMNS + 1)]

    with pytest.raises(CsvImportError) as err:
        parse_upload("excel", excel(wide))

    assert err.value.code == "CSV_PARSE_FAILED"


def test_a_header_below_the_scan_limit_is_not_found_in_a_workbook():
    rows = [["note"]] * (HEADER_SCAN_ROWS + 1) + TABLE

    with pytest.raises(CsvImportError) as err:
        parse_upload("excel", excel(rows))

    assert err.value.code == "CSV_MISSING_COLUMNS"


def test_a_file_that_is_not_a_workbook_gets_a_clear_error():
    with pytest.raises(CsvImportError) as err:
        read_workbook(b"this is not an excel file")

    assert err.value.code == "EXCEL_UNREADABLE"
    assert "password" in err.value.message  # protected workbooks are not zip files


def test_a_zip_that_is_not_a_workbook_gets_the_same_error():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("hello.txt", "hi")

    with pytest.raises(CsvImportError) as err:
        read_workbook(out.getvalue())

    assert err.value.code == "EXCEL_UNREADABLE"


def test_a_zip_bomb_is_refused_before_anything_is_unpacked():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"\0" * (MAX_WORKBOOK_UNCOMPRESSED + 1))
    assert len(out.getvalue()) < 100_000  # it really is small on the wire

    started = time.monotonic()
    with pytest.raises(CsvImportError) as err:
        read_workbook(out.getvalue())

    assert err.value.code == "EXCEL_TOO_LARGE"
    assert time.monotonic() - started < 2


def test_an_archive_with_thousands_of_parts_is_refused():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for i in range(2000):
            archive.writestr(f"p{i}.xml", "x")

    with pytest.raises(CsvImportError) as err:
        read_workbook(out.getvalue())

    assert err.value.code == "EXCEL_TOO_LARGE"


def test_xml_entity_protection_is_switched_on():
    # openpyxl only uses defusedxml when it is importable; this fails if that ever stops being true.
    assert openpyxl.DEFUSEDXML is True
    assert not openpyxl.LXML  # lxml would switch which parser reads some parts


def test_a_sheet_that_declares_xml_entities_is_refused():
    good = excel(TABLE)
    bomb = (
        b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY a "aaaaaaaaaa">'
        b'<!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">]>'
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&b;</t></is></c></row>'
        b"</sheetData></worksheet>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(good)) as source, zipfile.ZipFile(out, "w") as target:
        for item in source.infolist():
            data = bomb if item.filename == "xl/worksheets/sheet1.xml" else source.read(item)
            target.writestr(item.filename, data)

    with pytest.raises(CsvImportError) as err:
        read_workbook(out.getvalue())

    assert err.value.code == "EXCEL_UNREADABLE"


# ---------------------------------------------------------------- text encodings
def test_utf8_with_and_without_a_byte_order_mark_decodes():
    assert decode_text("Café".encode("utf-8-sig")) == "Café"
    assert decode_text("Café".encode()) == "Café"


def test_excels_windows_1252_csv_decodes_instead_of_turning_letters_into_marks():
    assert decode_text(b"Caf\xe9 au lait") == "Café au lait"


def test_excels_unicode_text_export_decodes_and_reads_as_a_tab_separated_table():
    text = "Date\tItem\tAmount\n01/08/2025\tCafé\t-5.00\n"

    preview = parse_upload("text", text.encode("utf-16"))  # utf-16 writes the byte order mark

    assert [(r.description, r.amount_sen) for r in preview.rows] == [("Café", 500)]


# ---------------------------------------------------------------- the downloadable template
# Keep identical to frontend/public/paysense-import-template.csv (the frontend has a test that the
# file starts with the same header). What the person downloads must import without a warning.
TEMPLATE_CSV = (
    "type,date,item,amount,category\n"
    "expense,2026-06-01,Nasi lemak,6.50,Food\n"
    "expense,2026-06-02,Bus to campus,3.00,Transport\n"
    "income,2026-06-05,PESP Allowance,1000.00,Allowance\n"
)


def test_the_template_the_person_downloads_imports_cleanly_as_it_stands():
    preview = parse_upload("text", TEMPLATE_CSV.encode())

    assert summary(preview) == [
        (date(2026, 6, 1), 650, "expense", "Food", "Nasi lemak"),
        (date(2026, 6, 2), 300, "expense", "Transport", "Bus to campus"),
        (date(2026, 6, 5), 100000, "income", "Allowance", "PESP Allowance"),
    ]
    assert not preview.skipped and all(not r.issues for r in preview.rows)


def test_the_template_also_imports_when_opened_and_saved_as_a_workbook():
    rows = [line.split(",") for line in TEMPLATE_CSV.strip().splitlines()]

    preview = parse_upload("excel", excel(rows))

    assert [r.category_guess for r in preview.rows] == ["Food", "Transport", "Allowance"]
